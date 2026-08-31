"""Renders a ReportPayload to .xlsx bytes via openpyxl — the same generic
shape the PDF renderer consumes, laid out as one worksheet."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.schemas.payload import ReportPayload


def render_excel(payload: ReportPayload) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = payload.title[:31] or "Report"  # Excel's own 31-char sheet-name limit

    row = 1
    sheet.cell(row=row, column=1, value=payload.title).font = Font(size=14, bold=True)
    row += 1
    if payload.company_name:
        sheet.cell(row=row, column=1, value=payload.company_name).font = Font(bold=True)
        row += 1
    sheet.cell(row=row, column=1, value=payload.period_label)
    row += 2

    for line in payload.summary:
        sheet.cell(row=row, column=1, value=line.label).font = Font(bold=True)
        sheet.cell(row=row, column=2, value=line.value)
        row += 1
    row += 1

    if payload.table is not None:
        for col, header in enumerate(payload.table.headers, start=1):
            cell = sheet.cell(row=row, column=col, value=header)
            cell.font = Font(bold=True)
        row += 1
        for line in payload.table.rows:
            for col, value in enumerate(line, start=1):
                sheet.cell(row=row, column=col, value=value)
            row += 1
        row += 1

    for note in payload.notes:
        cell = sheet.cell(row=row, column=1, value=f"Note: {note}")
        cell.font = Font(size=9, italic=True, color="808080")
        cell.alignment = Alignment(wrap_text=True)
        row += 1

    for col in range(1, 7):
        sheet.column_dimensions[get_column_letter(col)].width = 24

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
